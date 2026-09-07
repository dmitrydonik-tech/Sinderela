# ⚠ Пути внутри — абсолютные к build-сессии. Перед запуском замени базовые пути вверху под свой репозиторий (index.html / assets / временная папка). См. tools/README.md
# -*- coding: utf-8 -*-
import openpyxl, json
F="/sessions/awesome-ecstatic-tesla/mnt/SinderelaMD/price/Прейскурант.xlsx"
wb=openpyxl.load_workbook(F, data_only=True)
ws=wb.active
rows=[[ (None if x is None else str(x)) for x in r] for r in ws.iter_rows(values_only=True)]
json.dump(rows, open("/sessions/awesome-ecstatic-tesla/mnt/outputs/price_rows.json","w"), ensure_ascii=False)
print("sheet:", ws.title, "| rows:", len(rows), "| maxcol:", ws.max_column)
for r in rows[:8]:
    print(r)
